from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Iterable

import fitz  # PyMuPDF
import pdfplumber
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

from ..utils import compact_text, ensure_dir


def extract_native_pdf_text(path: Path, max_pages: int) -> str:
    """Extract embedded text from a text-based PDF.

    This is deliberately conservative. If this produces enough text, the pipeline
    can skip expensive OCR for text-native resumes.
    """
    chunks: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[:max_pages]:
                chunks.append(page.extract_text() or "")
    except Exception:
        return ""
    return compact_text("\n\n".join(chunks))


def _iter_docx_blocks(parent: DocxDocument | _Cell) -> Iterable[Paragraph | Table]:
    """Yield DOCX paragraphs and tables in document order."""
    if isinstance(parent, DocxDocument):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        return

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _paragraph_text(paragraph: Paragraph) -> str:
    # paragraph.text keeps visible run text and is enough for resume parsing.
    # Hyperlink visible text is also included in modern python-docx versions.
    return compact_text(paragraph.text)


def _table_to_markdown(table: Table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        values = [compact_text(cell.text.replace("\n", " ")) for cell in row.cells]
        if any(values):
            rows.append(values)

    if not rows:
        return ""

    # Most resume tables in the samples are two-column skill/category tables.
    # Markdown keeps key/value relationships clearer for the LLM parser.
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    separator = ["---"] * width
    body = normalized[1:]

    def line(values: list[str]) -> str:
        escaped = [value.replace("|", "\\|") for value in values]
        return "| " + " | ".join(escaped) + " |"

    return "\n".join([line(header), line(separator), *(line(row) for row in body)])


def extract_native_docx_text(path: Path) -> str:
    """Extract readable text from a .docx resume without OCR.

    DOCX resumes are text-native. This function preserves normal paragraphs,
    table content, headers, and footers so resumes that keep skills in tables are
    parsed correctly by the existing LLM extraction step.
    """
    try:
        document = Document(path)
    except Exception:
        return ""

    chunks: list[str] = []

    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = _paragraph_text(block)
            if text:
                chunks.append(text)
        elif isinstance(block, Table):
            text = _table_to_markdown(block)
            if text:
                chunks.append(text)

    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                text = _paragraph_text(paragraph)
                if text:
                    chunks.append(text)
            for table in container.tables:
                text = _table_to_markdown(table)
                if text:
                    chunks.append(text)

    return compact_text("\n\n".join(chunks))


def render_pdf_pages_to_images(
    path: Path,
    output_dir: Path,
    *,
    max_pages: int,
    dpi: int,
    max_dimension: int,
    jpeg_quality: int,
) -> list[Path]:
    ensure_dir(output_dir)
    image_paths: list[Path] = []
    doc = fitz.open(path)
    try:
        page_count = min(len(doc), max_pages)
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)
        for index in range(page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            image_path = output_dir / f"page_{index + 1:03d}.jpg"
            pix.save(str(image_path))
            optimize_image(image_path, image_path, max_dimension=max_dimension, jpeg_quality=jpeg_quality)
            image_paths.append(image_path)
    finally:
        doc.close()
    return image_paths


def optimize_image(input_path: Path, output_path: Path, *, max_dimension: int, jpeg_quality: int) -> Path:
    with Image.open(input_path) as img:
        img = ImageOps.exif_transpose(img)
        img = img.convert("RGB")
        if max(img.size) > max_dimension:
            img.thumbnail((max_dimension, max_dimension), Image.Resampling.LANCZOS)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        img.save(output_path, format="JPEG", quality=jpeg_quality, optimize=True)
    return output_path


def image_to_optimized_jpeg(
    path: Path,
    output_dir: Path,
    *,
    max_dimension: int,
    jpeg_quality: int,
) -> Path:
    ensure_dir(output_dir)
    output_path = output_dir / f"{path.stem}.jpg"
    return optimize_image(path, output_path, max_dimension=max_dimension, jpeg_quality=jpeg_quality)


class TempWorkDir:
    def __init__(self, prefix: str = "resume_ocr_"):
        self.prefix = prefix
        self._tmp: tempfile.TemporaryDirectory[str] | None = None
        self.path: Path | None = None

    def __enter__(self) -> Path:
        self._tmp = tempfile.TemporaryDirectory(prefix=self.prefix)
        self.path = Path(self._tmp.name)
        return self.path

    def __exit__(self, exc_type, exc, tb) -> None:
        if self._tmp is not None:
            self._tmp.cleanup()
